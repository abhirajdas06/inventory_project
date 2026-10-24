"""Stock Out Register — for auditors.

Every stock-out transaction for one date (or a date range): what went out,
to which client, on which invoice / OLF-DC, with what status, from which
store, by whom — and what has happened to it since (still out, returned …).
Filterable by status, store and free text; exportable to a flat Excel sheet
(one row per stock out, real date cells, auto-filter) that works directly
with pivot tables.

Each stock out is its own row, so an item that went out, came back and went
out again appears once per stock out — that is what an audit needs.
"""
from datetime import date, datetime
from io import BytesIO

from django.core.paginator import EmptyPage, PageNotAnInteger, Paginator
from django.db.models import Count, Exists, OuterRef, Q, Subquery
from django.db.models.functions import Coalesce, TruncDate
from django.http import HttpResponse
from django.shortcuts import render
from django.utils import timezone
from openpyxl.utils import get_column_letter

from apps.core.activity import log_activity
from apps.core.permissions import has_permission, require_permission
from apps.inventory.models import InventoryTransaction, RentalRecord

PER_PAGE = 100

# Statuses a stock out can carry (everything except the in-stock LIVE).
OUT_STATUSES = [(k, l) for k, l in InventoryTransaction.STOCK_STATUS if k != 'LIVE']

_CATEGORY_ATTRS = ('spare', 'card', 'cpu', 'controller', 'memory', 'sfp', 'railkit', 'harddisk', 'networking_spare', 'server')

COLUMNS = [
    # (heading, row key, Excel column width)
    ('Stock Out Date', 'out_date', 14),
    ('Status', 'status_label', 14),
    ('Category', 'category', 20),
    ('Brand', 'brand', 14),
    ('Model', 'model', 20),
    ('Part No', 'part_no', 18),
    ('Serial No', 'serial_no', 22),
    ('Barcode No', 'barcode', 16),
    ('Client', 'client', 26),
    ('Invoice No', 'invoice', 16),
    ('OLF / DC No', 'olf_dc', 16),
    ('Store', 'store', 10),
    ('Location', 'location', 28),
    ('Part Of', 'installed_in', 30),
    ('Expected Return (Rent)', 'expected_return', 14),
    ('Stocked Out By', 'performed_by', 16),
    ('Entered On', 'entered_on', 18),
    ('Current State', 'current_state', 26),
]


def _parse_date(value):
    try:
        return datetime.strptime(str(value or '').strip(), '%Y-%m-%d').date()
    except ValueError:
        return None


def read_filters(request):
    """Normalised filters from the query string (shared by page and export)."""
    get = request.GET
    today = timezone.localdate()
    mode = 'range' if get.get('mode') == 'range' else 'day'
    if mode == 'day':
        day = _parse_date(get.get('date')) or today
        date_from = date_to = day
    else:
        date_from = _parse_date(get.get('date_from')) or today
        date_to = _parse_date(get.get('date_to')) or date_from
        if date_from > date_to:
            date_from, date_to = date_to, date_from
    status = (get.get('status') or '').strip().upper()
    if status not in dict(OUT_STATUSES):
        status = ''
    store = (get.get('store') or '').strip().upper()
    if store not in dict(InventoryTransaction.STORE_LOCATION):
        store = ''
    return {
        'mode': mode,
        'date': date_from,
        'date_from': date_from,
        'date_to': date_to,
        'status': status,
        'store': store,
        'q': ' '.join((get.get('q') or '').split()),
        'main_only': get.get('main_only') == '1',
    }


def register_queryset(filters, with_status=True):
    """All stock-out transactions matching the filters, oldest first."""
    from apps.categories.models import Spare
    from apps.servers.models import ServerComponent

    latest = InventoryTransaction.objects.filter(product=OuterRef('product')).order_by('-created_at', '-id')
    qs = (
        InventoryTransaction.objects.filter(transaction_type='OUT')
        # Older imported rows may have no stock_out_date — fall back to the day it was entered.
        .annotate(out_date=Coalesce('stock_out_date', TruncDate('created_at')))
        .filter(out_date__gte=filters['date_from'], out_date__lte=filters['date_to'])
        .annotate(
            in_server=Exists(ServerComponent.objects.filter(product=OuterRef('product'))),
            in_controller=Exists(Spare.objects.filter(product=OuterRef('product'), controller__isnull=False)),
        )
    )
    if with_status and filters['status']:
        qs = qs.filter(stock_status=filters['status'])
    if filters['store']:
        qs = qs.filter(store_location=filters['store'])
    if filters['main_only']:
        qs = qs.filter(in_server=False, in_controller=False)
    q = filters['q']
    if q:
        barcode_q = Q()
        for attr in _CATEGORY_ATTRS:
            barcode_q |= Q(**{f'product__{attr}__barcode__icontains': q})
        qs = qs.filter(
            Q(client_name__icontains=q) | Q(invoice_no__icontains=q) | Q(olf_dc_number__icontains=q)
            | Q(product__serial_no__icontains=q) | Q(product__part_no__icontains=q)
            | Q(product__name__icontains=q) | Q(product__category__name__icontains=q)
            | Q(product__server_component__barcode__icontains=q) | barcode_q
        ).distinct()
    return qs.annotate(
        latest_id=Subquery(latest.values('id')[:1]),
        latest_type=Subquery(latest.values('transaction_type')[:1]),
        latest_status=Subquery(latest.values('stock_status')[:1]),
        latest_at=Subquery(latest.values('created_at')[:1]),
        expected_return=Subquery(
            RentalRecord.objects.filter(product=OuterRef('product'), rent_out_date=OuterRef('out_date'))
            .order_by('-id').values('expected_return_date')[:1]
        ),
    ).order_by('out_date', 'created_at', 'id')


def _with_details(qs):
    related = ['product__category', 'performed_by']
    for attr in _CATEGORY_ATTRS:
        related += [f'product__{attr}', f'product__{attr}__brand']
    related.append('product__spare__controller__product')
    return qs.select_related(*related).prefetch_related('product__server_component__server')


def _category_object(product):
    for attr in _CATEGORY_ATTRS:
        try:
            obj = getattr(product, attr)
        except Exception:
            obj = None
        if obj is not None:
            return obj
    return None


def build_row(txn):
    product = txn.product
    obj = _category_object(product)
    component = next(iter(product.server_component.all()), None)
    brand = getattr(obj, 'brand', None)

    installed_in = ''
    if component is not None and component.server is not None:
        s = component.server
        installed_in = f"Server: {s.machine_no or s.model or ''} [{s.service_tag}]"
    else:
        try:
            spare = product.spare
        except Exception:
            spare = None
        if spare is not None and spare.controller_id:
            c = spare.controller
            installed_in = f"Controller: {c.model or ''} [{getattr(c.product, 'serial_no', '') or ''}]"

    if txn.latest_id == txn.id:
        current = 'Still out'
    elif txn.latest_type == 'IN':
        current = f"Back in stock ({txn.latest_status})"
    elif txn.latest_type == 'OUT':
        current = f"Stocked out again ({txn.latest_status})"
    else:
        current = f"{txn.latest_type or ''} {txn.latest_status or ''}".strip()
    if txn.latest_id != txn.id and txn.latest_at:
        current += f" on {timezone.localtime(txn.latest_at):%d-%m-%Y}"

    return {
        'id': txn.id,
        'product_id': product.id,
        'out_date': txn.out_date,
        'status': txn.stock_status,
        'status_label': dict(InventoryTransaction.STOCK_STATUS).get(txn.stock_status, txn.stock_status),
        'category': (product.category.name if product.category_id else '') or (component.spare_type if component else ''),
        'brand': (brand.name if brand else '') or product.brand or '',
        'model': getattr(obj, 'model', None) or getattr(obj, 'size', None) or product.model or '',
        'part_no': getattr(obj, 'part_no', None) or (component.part_no if component else '') or product.part_no or '',
        'serial_no': product.serial_no or (component.serial_no if component else '') or '',
        'barcode': getattr(obj, 'barcode', None) or (component.barcode if component else '') or '',
        'client': txn.client_name or '',
        'invoice': txn.invoice_no or '',
        'olf_dc': txn.olf_dc_number or '',
        'store': txn.store_location or '',
        'location': getattr(obj, 'location', None) or (component.location if component else '') or '',
        'installed_in': installed_in,
        'expected_return': txn.expected_return if txn.stock_status == 'RENT' else None,
        'performed_by': txn.performed_by.get_username() if txn.performed_by else '',
        'entered_on': timezone.localtime(txn.created_at) if txn.created_at else None,
        'current_state': current,
        'still_out': txn.latest_id == txn.id,
    }


def _status_counts(filters):
    rows = (register_queryset(filters, with_status=False).order_by()
            .values('stock_status').annotate(n=Count('id', distinct=True)))
    counts = {r['stock_status']: r['n'] for r in rows}
    return [(k, l, counts[k]) for k, l in OUT_STATUSES if counts.get(k)]


@require_permission('stock_out_register')
def stock_out_register(request):
    filters = read_filters(request)
    qs = register_queryset(filters)
    paginator = Paginator(_with_details(qs), PER_PAGE)
    try:
        page_obj = paginator.page(request.GET.get('page', 1))
    except (PageNotAnInteger, EmptyPage):
        page_obj = paginator.page(1)
    status_counts = _status_counts(filters)
    totals = qs.order_by().aggregate(
        clients=Count('client_name', distinct=True, filter=~Q(client_name='') & Q(client_name__isnull=False)),
        invoices=Count('invoice_no', distinct=True, filter=~Q(invoice_no='') & Q(invoice_no__isnull=False)),
    )
    return render(request, 'inventory/stock_out_register.html', {
        'filters': filters,
        'rows': [build_row(t) for t in page_obj.object_list],
        'page_obj': page_obj,
        'total_count': paginator.count,
        'status_counts': status_counts,
        'all_status_total': sum(n for _, _, n in status_counts),
        'unique_clients': totals['clients'],
        'unique_invoices': totals['invoices'],
        'status_options': OUT_STATUSES,
        'store_options': InventoryTransaction.STORE_LOCATION,
        'can_history': has_permission(request.user, 'product_history'),
        'export_query': request.GET.urlencode(),
    })


@require_permission('stock_out_register')
def stock_out_register_export(request):
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill

    filters = read_filters(request)
    qs = _with_details(register_queryset(filters))
    book = Workbook(write_only=True)
    sheet = book.create_sheet('Stock Out Register')
    sheet.freeze_panes = 'A2'
    for index, (_, _, width) in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    bold, fill = Font(bold=True), PatternFill(fill_type='solid', fgColor='DCEBFF')
    header = []
    for heading, _, _ in COLUMNS:
        cell = WriteOnlyCell(sheet, value=heading)
        cell.font, cell.fill = bold, fill
        header.append(cell)
    sheet.append(header)

    count = 0
    for txn in qs:  # not .iterator(): it would skip the prefetch
        row = build_row(txn)
        values = []
        for _, key, _ in COLUMNS:
            value = row[key]
            if isinstance(value, datetime):
                value = value.replace(tzinfo=None)
            values.append('' if value is None else value)
        sheet.append(values)
        count += 1
    last_col = get_column_letter(len(COLUMNS))
    sheet.auto_filter.ref = f'A1:{last_col}{count + 1}'

    info = book.create_sheet('Filters')
    period = (f"{filters['date_from']:%d-%m-%Y}" if filters['mode'] == 'day'
              else f"{filters['date_from']:%d-%m-%Y} to {filters['date_to']:%d-%m-%Y}")
    for line in (
        ['Stock Out Register'],
        ['Period', period],
        ['Status', dict(OUT_STATUSES).get(filters['status'], 'All')],
        ['Store', filters['store'] or 'All'],
        ['Search', filters['q'] or '—'],
        ['Parts inside servers / controllers', 'Hidden' if filters['main_only'] else 'Included'],
        ['Rows', count],
        ['Exported by', request.user.get_username()],
        ['Exported at', timezone.localtime().replace(tzinfo=None)],
    ):
        info.append(line)

    stream = BytesIO()
    book.save(stream)
    log_activity(action='STOCK_OUT_REGISTER_EXPORT', module='EXPORT', entity='Stock Out Register',
                 user=request.user, new_values={**{k: str(v) for k, v in filters.items()}, 'rows': count},
                 remarks=f'Stock Out Register exported ({period}, {count} rows)')
    suffix = (f"{filters['date_from']}" if filters['mode'] == 'day'
              else f"{filters['date_from']}_to_{filters['date_to']}")
    response = HttpResponse(stream.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="stock-out-register_{suffix}.xlsx"'
    return response
