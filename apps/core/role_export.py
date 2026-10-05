"""Excel of every user type, what it may do, and who has it — for management
to review and sign off which role each user should get.

Sheets:
  1. Role Permissions — permission x user type matrix (Yes / blank)
  2. User Types       — one row per type: users, permissions, approval columns
  3. Users            — each user's current type + "Approved User Type" to fill in
"""
from django.contrib.auth.models import User
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from apps.core.models import RolePermission
from apps.core.permissions import PERMISSION_LABELS, ROLE_PERMISSIONS, all_roles

HEADER_FILL = PatternFill(fill_type='solid', fgColor='DCEBFF')
YES_FILL = PatternFill(fill_type='solid', fgColor='D9F2E3')
FILL_IN_FILL = PatternFill(fill_type='solid', fgColor='FFF7D6')
THIN = Side(style='thin', color='C9D3E0')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='top')
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)


def role_permission_table():
    """[(role, label, kind, permissions set)] as they are right now."""
    rows = {row.role: row for row in RolePermission.objects.all()}
    table = []
    for role, label in all_roles():
        row = rows.get(role)
        if row is not None:
            perms = set(row.permissions or [])
            kind = 'Added type' if row.is_custom else 'Built-in (customised)'
        else:
            perms = set(ROLE_PERMISSIONS.get(role, set()))
            kind = 'Built-in (default)'
        if role == 'ADMIN':
            perms.add('user_management')  # Admin can never lose this
        table.append((role, label, kind, perms))
    return table


def _user_role(user):
    profile = getattr(user, 'profile', None)
    if profile is not None:
        return profile.role
    return 'ADMIN' if user.is_superuser else ''


def _header(sheet, headers, row=1):
    for col, text in enumerate(headers, start=1):
        cell = sheet.cell(row=row, column=col, value=text)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = CENTER


def _title(sheet, text, generated_by, width):
    sheet.cell(row=1, column=1, value=text).font = Font(bold=True, size=14)
    stamp = timezone.localtime().strftime('%d-%m-%Y %H:%M')
    sheet.cell(row=2, column=1, value=f'Generated {stamp} by {generated_by} from the live system.').font = Font(italic=True, color='667085')
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(width, 1))
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(width, 1))


def _sign_off(sheet, start_row):
    labels = ['Reviewed by (Management):', 'Signature:', 'Date:', 'Comments:']
    for offset, text in enumerate(labels):
        cell = sheet.cell(row=start_row + offset, column=1, value=text)
        cell.font = Font(bold=True)
        sheet.cell(row=start_row + offset, column=2).border = Border(bottom=THIN)
        sheet.cell(row=start_row + offset, column=3).border = Border(bottom=THIN)


def build_role_approval_workbook(generated_by=''):
    table = role_permission_table()
    users = list(User.objects.select_related('profile').order_by('username'))
    label_by_role = {role: label for role, label, _, _ in table}
    user_count = {}
    for user in users:
        role = _user_role(user)
        user_count[role] = user_count.get(role, 0) + 1

    book = Workbook()

    # ── 1. Permission matrix ──────────────────────────────────────────────
    matrix = book.active
    matrix.title = 'Role Permissions'
    width = len(table) + 1
    _title(matrix, 'User Types and Permissions — for Management Approval', generated_by, width)
    _header(matrix, ['Permission'] + [label for _, label, _, _ in table], row=4)
    for r, (perm, perm_label) in enumerate(PERMISSION_LABELS.items(), start=5):
        cell = matrix.cell(row=r, column=1, value=perm_label)
        cell.border = BORDER
        cell.alignment = WRAP
        for c, (_, _, _, perms) in enumerate(table, start=2):
            has = perm in perms
            cell = matrix.cell(row=r, column=c, value='Yes' if has else '')
            cell.border = BORDER
            cell.alignment = CENTER
            if has:
                cell.fill = YES_FILL
                cell.font = Font(bold=True, color='1E7B45')
    total_row = 5 + len(PERMISSION_LABELS)
    for label, values in (('Users with this type', [user_count.get(role, 0) for role, _, _, _ in table]),
                          ('Type', [kind for _, _, kind, _ in table])):
        matrix.cell(row=total_row, column=1, value=label).font = Font(bold=True)
        matrix.cell(row=total_row, column=1).border = BORDER
        for c, value in enumerate(values, start=2):
            cell = matrix.cell(row=total_row, column=c, value=value)
            cell.border = BORDER
            cell.alignment = CENTER
        total_row += 1
    matrix.column_dimensions['A'].width = 52
    for c in range(2, width + 1):
        matrix.column_dimensions[get_column_letter(c)].width = 16
    matrix.row_dimensions[4].height = 34
    matrix.freeze_panes = 'B5'
    _sign_off(matrix, total_row + 2)

    # ── 2. One row per user type ──────────────────────────────────────────
    types = book.create_sheet('User Types')
    headers = ['User Type', 'Type', 'No. of Users', 'Users', 'Permissions', 'Approved (Yes/No)', 'Management Remark']
    _title(types, 'User Types — Approve Each Type', generated_by, len(headers))
    _header(types, headers, row=4)
    users_by_role = {}
    for user in users:
        users_by_role.setdefault(_user_role(user), []).append(user.username)
    for r, (role, label, kind, perms) in enumerate(table, start=5):
        values = [
            label, kind, user_count.get(role, 0), ', '.join(users_by_role.get(role, [])) or '—',
            '\n'.join(f'• {PERMISSION_LABELS[p]}' for p in PERMISSION_LABELS if p in perms) or 'No permissions',
            '', '',
        ]
        for c, value in enumerate(values, start=1):
            cell = types.cell(row=r, column=c, value=value)
            cell.border = BORDER
            cell.alignment = WRAP
            if c >= 6:
                cell.fill = FILL_IN_FILL
    for col, w in zip('ABCDEFG', (24, 20, 12, 30, 60, 18, 34)):
        types.column_dimensions[col].width = w
    types.freeze_panes = 'B5'
    last = 4 + len(table)
    yes_no = DataValidation(type='list', formula1='"Yes,No"', allow_blank=True)
    types.add_data_validation(yes_no)
    if table:
        yes_no.add(f'F5:F{last}')
    _sign_off(types, last + 3)

    # ── 3. Users and the type each should get ─────────────────────────────
    people = book.create_sheet('Users')
    headers = ['Username', 'Email', 'Active', 'Current User Type', 'Approved User Type', 'Management Remark']
    _title(people, 'Users — Confirm the User Type for Each Person', generated_by, len(headers))
    _header(people, headers, row=4)
    for r, user in enumerate(users, start=5):
        role = _user_role(user)
        values = [user.username, user.email or '', 'Yes' if user.is_active else 'No',
                  label_by_role.get(role, role or '(none)'), '', '']
        for c, value in enumerate(values, start=1):
            cell = people.cell(row=r, column=c, value=value)
            cell.border = BORDER
            cell.alignment = WRAP
            if c >= 5:
                cell.fill = FILL_IN_FILL
    for col, w in zip('ABCDEF', (22, 30, 9, 26, 26, 36)):
        people.column_dimensions[col].width = w
    people.freeze_panes = 'A5'
    last = 4 + len(users)
    # Dropdown of the existing user types, listed on a hidden helper sheet.
    lists = book.create_sheet('Lists')
    for r, (_, label, _, _) in enumerate(table, start=1):
        lists.cell(row=r, column=1, value=label)
    lists.sheet_state = 'hidden'
    if users and table:
        pick = DataValidation(type='list', formula1=f"=Lists!$A$1:$A${len(table)}", allow_blank=True)
        people.add_data_validation(pick)
        pick.add(f'E5:E{last}')
    _sign_off(people, last + 3)

    return book
